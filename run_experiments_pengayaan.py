"""Pengujian fitur pengayaan: m-bit LSB, steganografi audio WAV, steganalisis chi-square.

Jalankan:  python run_experiments_pengayaan.py
Keluaran (folder hasil_uji/):
  hasil_pengayaan.xlsx           tabel semua pengujian pengayaan
  mbit_tradeoff.png              grafik trade-off kapasitas vs PSNR (m = 1..6)
  chi_square_pvalue.png          p-value chi-square vs tingkat pengisian
  chi_square_selisih_pasangan.png  selisih h[2i]-h[2i+1], cover vs stego
  audio_snr.png                  SNR audio vs ukuran pesan & m
  (sample_audio/*.wav dibuat otomatis bila belum ada)
"""
import os
import time
import wave

import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from stego import core, crypto, metrics, audio, steganalysis

BASE = os.path.dirname(os.path.abspath(__file__))
IMG_DIR = os.path.join(BASE, "sample_images")
AUD_DIR = os.path.join(BASE, "sample_audio")
OUT_DIR = os.path.join(BASE, "hasil_uji")

IMAGES = ["01_gradien", "02_derau", "03_geometri", "04_foto_sintetis", "05_papan_teks"]
PASSWORD = "sandi-uji-pengayaan"     # hanya untuk pengujian otomatis
STEGO_KEY = "kunci-uji-pengayaan"
OVERHEAD = crypto.SALT_LEN + crypto.NONCE_LEN + 16     # salt + nonce + tag GCM = 44 B
GAGAL = (core.ExtractionError, crypto.DecryptionError)


def rand_msg(n, seed=7):
    return np.random.default_rng(seed).bytes(n)


def load(name):
    return core.load_image_rgb(os.path.join(IMG_DIR, name + ".png"))


def write_sheet(wb, title, header, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="3465E0")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    for i, col in enumerate(ws.columns):
        w = (widths[i] if widths else None) or max(len(str(c.value or "")) for c in col) + 3
        ws.column_dimensions[col[0].column_letter].width = min(max(w, 10), 46)
    ws.freeze_panes = "A2"
    return ws


# ---------------------------------------------------------------------------
# 1) m-bit LSB: trade-off kapasitas vs PSNR
# ---------------------------------------------------------------------------
def uji_mbit():
    print("[1/3] m-bit LSB ...")
    rows_fixed, rows_full = [], []
    for name in IMAGES:
        img = load(name)
        for m in range(1, 7):
            cap = core.capacity_bytes(img, m)
            for label, n, dest in (("tetap 5 KB", 5120, rows_fixed),
                                   ("penuh (kapasitas maks.)", cap - OVERHEAD, rows_full)):
                msg = rand_msg(n)
                st = core.embed(img, msg, PASSWORD, STEGO_KEY, m)
                ok = core.extract(st, PASSWORD, STEGO_KEY) == msg
                ps = metrics.psnr(img, st)
                dest.append([name + ".png", m, label, n, cap,
                             round(metrics.mse(img, st), 5), round(ps, 2),
                             round(metrics.changed_pixels_percent(img, st), 3),
                             "YA" if ps >= 30 else "TIDAK", "YA" if ok else "TIDAK"])
        print("   ", name)
    hdr = ["Citra", "m (bit/kanal)", "Skenario", "Pesan (byte)", "Kapasitas (byte)",
           "MSE", "PSNR (dB)", "Kanal Berubah (%)", "PSNR >= 30 dB", "Ekstraksi Utuh"]

    # grafik
    fig, ax = plt.subplots(1, 3, figsize=(15, 4.4))
    ms = np.arange(1, 7)
    theory = 10 * np.log10(255 ** 2 / ((4.0 ** ms - 1) / 6))
    for name in IMAGES:
        full = [r for r in rows_full if r[0] == name + ".png"]
        fixed = [r for r in rows_fixed if r[0] == name + ".png"]
        ax[0].plot(ms, [r[6] for r in full], "o-", label=name)
        ax[1].plot(ms, [r[6] for r in fixed], "o-", label=name)
    ax[0].plot(ms, theory, "k--", lw=2, label="teori (bit acak)")
    for a, t in ((ax[0], "Pesan memenuhi kapasitas maks."), (ax[1], "Pesan tetap 5 KB")):
        a.axhline(30, color="red", ls=":", label="ambang 30 dB")
        a.set_title(t); a.set_xlabel("m (bit per kanal)"); a.set_ylabel("PSNR (dB)")
        a.grid(alpha=.3)
    ax[0].legend(fontsize=7)
    cap512 = [core.capacity_bytes(load(IMAGES[0]), m) / 1024 for m in ms]
    ax[2].bar(ms, cap512, color="#3465e0")
    for x, v in zip(ms, cap512):
        ax[2].text(x, v, f"{v:.0f}", ha="center", va="bottom", fontsize=8)
    ax[2].set_title("Kapasitas citra 512x512 (KB)")
    ax[2].set_xlabel("m (bit per kanal)"); ax[2].set_ylabel("KB")
    fig.suptitle("Trade-off kapasitas vs kualitas pada m-bit LSB", fontweight="bold")
    fig.tight_layout()
    fig.savefig(os.path.join(OUT_DIR, "mbit_tradeoff.png"), dpi=130)
    plt.close(fig)
    return hdr, rows_fixed, rows_full


# ---------------------------------------------------------------------------
# 2) Steganalisis chi-square
# ---------------------------------------------------------------------------
def uji_chisq():
    print("[2/3] chi-square ...")
    fills = [0.0, 0.1, 0.25, 0.5, 0.75, 1.0]
    rows, curves = [], {}
    stego_full = {}
    for name in IMAGES:
        img = load(name)
        cap = core.capacity_bytes(img, 1)
        curves[name] = []
        for f in fills:
            if f == 0.0:
                st, n = img, 0
            else:
                n = int((cap - OVERHEAD) * f)
                st = core.embed(img, rand_msg(n), PASSWORD, STEGO_KEY, 1)
            r = steganalysis.analyze(st)
            rate = (n + OVERHEAD + core.HEADER_LEN) * 8 / img.size * 100
            rows.append([name + ".png", "cover (tanpa pesan)" if f == 0 else f"{int(f * 100)}% kapasitas",
                         n, round(rate, 2), round(r["chi2"], 2), r["df"],
                         float(f"{r['p_value']:.4g}"), "YA" if r["flagged"] else "TIDAK",
                         "benar" if (bool(r["flagged"]) == (f > 0)) else "SALAH"])
            curves[name].append(max(r["p_value"], 1e-300))
            if f == 1.0:
                stego_full[name] = st
        print("   ", name)
    hdr = ["Citra", "Skenario", "Pesan (byte)", "Laju sisip (% slot)", "chi-square", "df",
           "p-value", "Terdeteksi (p > 0.05)", "Keputusan"]

    fig, ax = plt.subplots(figsize=(8, 4.8))
    for name in IMAGES:
        ax.plot([f * 100 for f in fills], curves[name], "o-", label=name)
    ax.set_yscale("log"); ax.axhline(0.05, color="red", ls=":", label="alpha = 0.05")
    ax.set_xlabel("Pengisian kapasitas (%)"); ax.set_ylabel("p-value (log)")
    ax.set_title("Uji chi-square vs tingkat pengisian (LSB 1-bit, posisi acak)")
    ax.grid(alpha=.3); ax.legend(fontsize=8)
    fig.tight_layout(); fig.savefig(os.path.join(OUT_DIR, "chi_square_pvalue.png"), dpi=130)
    plt.close(fig)

    name = "04_foto_sintetis"
    cov, st = load(name), stego_full[name]
    fig, ax = plt.subplots(1, 2, figsize=(11, 3.8), sharey=True)
    ax[0].bar(range(128), steganalysis.pair_difference(cov), width=1, color="#3465e0")
    ax[0].set_title("Cover: h[2i] - h[2i+1]")
    ax[1].bar(range(128), steganalysis.pair_difference(st), width=1, color="#d1373f")
    ax[1].set_title("Stego 100% kapasitas: mendekati 0")
    for a in ax:
        a.set_xlabel("indeks pasangan i")
    ax[0].set_ylabel("selisih frekuensi")
    fig.suptitle(f"Penyamarataan pasangan nilai ({name})", fontweight="bold")
    fig.tight_layout(); fig.savefig(os.path.join(OUT_DIR, "chi_square_selisih_pasangan.png"), dpi=130)
    plt.close(fig)
    return hdr, rows


# ---------------------------------------------------------------------------
# 3) Audio WAV
# ---------------------------------------------------------------------------
def buat_sampel_audio():
    """Buat 3 berkas WAV sintetis (deterministik) bila belum ada."""
    os.makedirs(AUD_DIR, exist_ok=True)
    rate, dur = 22050, 3.0
    t = np.arange(int(rate * dur)) / rate
    rng = np.random.default_rng(42)
    specs = {}
    chord = sum(np.sin(2 * np.pi * f * t) for f in (261.6, 329.6, 392.0)) / 3
    chord = chord * (0.6 + 0.4 * np.sin(2 * np.pi * 0.7 * t)) + 0.01 * rng.standard_normal(t.size)
    specs["01_nada_16bit_mono.wav"] = (2, [chord])
    sweep = np.sin(2 * np.pi * (200 * t + 400 * t ** 2))
    left = 0.5 * sweep + 0.05 * rng.standard_normal(t.size)
    right = 0.5 * np.sin(2 * np.pi * 330 * t) * np.exp(-t / 2) + 0.05 * rng.standard_normal(t.size)
    specs["02_sapuan_16bit_stereo.wav"] = (2, [left, right])
    tone = np.sin(2 * np.pi * 440 * t) * (0.5 + 0.5 * np.sin(2 * np.pi * 3 * t)) + 0.02 * rng.standard_normal(t.size)
    specs["03_tremolo_8bit_mono.wav"] = (1, [tone])
    for fname, (width, chans) in specs.items():
        path = os.path.join(AUD_DIR, fname)
        if os.path.exists(path):
            continue
        x = np.clip(np.stack(chans, axis=1) * 0.9, -1, 1).reshape(-1)
        if width == 2:
            raw = (x * 32767).astype("<i2").tobytes()
        else:
            raw = ((x * 127) + 128).astype(np.uint8).tobytes()
        with wave.open(path, "wb") as w:
            w.setnchannels(len(chans)); w.setsampwidth(width)
            w.setframerate(rate); w.writeframes(raw)
    return sorted(specs)


def uji_audio():
    print("[3/3] audio WAV ...")
    files = buat_sampel_audio()
    rows, snr_table = [], {}
    for fname in files:
        path = os.path.join(AUD_DIR, fname)
        x, params = audio.load_wav(path)
        xs = audio.to_signed(x, params["width"])
        peak = audio.peak_value(params["width"])
        for m in (1, 2, 3):
            cap = audio.capacity_bytes(x, m)
            for label, n in (("kecil (100 B)", 100), ("sedang (1 KB)", 1024), ("besar (5 KB)", 5120)):
                if n + OVERHEAD > cap:
                    rows.append([fname, params["width"] * 8, params["channels"], m, label, n, cap,
                                 "-", "-", "-", "DITOLAK (melebihi kapasitas)"])
                    continue
                msg = rand_msg(n)
                st = audio.embed(x, msg, PASSWORD, STEGO_KEY, m)
                ok = audio.extract(st, PASSWORD, STEGO_KEY) == msg
                ys = audio.to_signed(st, params["width"])
                sn, ps, ms_ = metrics.snr(xs, ys), metrics.psnr(xs, ys, peak), metrics.mse(xs, ys)
                rows.append([fname, params["width"] * 8, params["channels"], m, label, n, cap,
                             round(ms_, 4), round(sn, 2), round(ps, 2),
                             "UTUH" if ok else "RUSAK"])
                snr_table.setdefault((fname, m), []).append((n, sn))
        # uji kunci salah + kerapuhan (simpan ulang bukan 16->8 bit, cukup uji kunci salah)
        st = audio.embed(x, rand_msg(1024), PASSWORD, STEGO_KEY, 1)
        try:
            audio.extract(st, PASSWORD, "kunci-salah")
            hasil = "BERHASIL (tidak diharapkan)"
        except GAGAL:
            hasil = "GAGAL (sesuai harapan)"
        rows.append([fname, params["width"] * 8, params["channels"], 1, "uji stego-key salah",
                     1024, audio.capacity_bytes(x, 1), "-", "-", "-", hasil])
        print("   ", fname)
    hdr = ["Berkas WAV", "Bit", "Kanal", "m", "Ukuran pesan", "Pesan (byte)", "Kapasitas (byte)",
           "MSE", "SNR (dB)", "PSNR (dB)", "Ekstraksi / Hasil"]

    fig, ax = plt.subplots(figsize=(8, 4.6))
    for (fname, m), pts in sorted(snr_table.items()):
        ax.plot([p[0] for p in pts], [p[1] for p in pts], "o-",
                label=f"{fname[:2]} m={m}")
    ax.axhline(30, color="red", ls=":", label="30 dB")
    ax.set_xscale("log"); ax.set_xlabel("Ukuran pesan (byte)"); ax.set_ylabel("SNR (dB)")
    ax.set_title("SNR audio stego vs ukuran pesan (01/02: 16-bit, 03: 8-bit)")
    ax.grid(alpha=.3); ax.legend(fontsize=7, ncol=2)
    fig.tight_layout(); fig.savefig(os.path.join(OUT_DIR, "audio_snr.png"), dpi=130)
    plt.close(fig)
    return hdr, rows


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    t0 = time.time()
    hdr_m, fixed, full = uji_mbit()
    hdr_c, chi = uji_chisq()
    hdr_a, aud = uji_audio()

    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "mbit - pesan tetap 5KB", hdr_m, fixed)
    write_sheet(wb, "mbit - kapasitas penuh", hdr_m, full)
    write_sheet(wb, "Chi-Square", hdr_c, chi)
    write_sheet(wb, "Audio WAV", hdr_a, aud)
    wb.save(os.path.join(OUT_DIR, "hasil_pengayaan.xlsx"))
    print(f"Selesai dalam {time.time() - t0:.0f} detik -> {OUT_DIR}")


if __name__ == "__main__":
    main()
