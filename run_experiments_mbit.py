"""Pengujian fitur pengayaan: varian m-bit LSB (trade-off kapasitas vs PSNR).

Jalankan:  python run_experiments_mbit.py
Keluaran (folder hasil_uji/):
  hasil_mbit.xlsx     tabel PSNR/MSE/kapasitas untuk m = 1..6
  mbit_tradeoff.png   grafik trade-off kapasitas vs PSNR
"""
import os
import time

import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from stego import core, crypto, metrics

BASE = os.path.dirname(os.path.abspath(__file__))
IMG_DIR = os.path.join(BASE, "sample_images")
OUT_DIR = os.path.join(BASE, "hasil_uji")

IMAGES = ["01_gradien", "02_derau", "03_geometri", "04_foto_sintetis", "05_papan_teks"]
PASSWORD = "sandi-uji-pengayaan"     # hanya untuk pengujian otomatis
STEGO_KEY = "kunci-uji-pengayaan"
OVERHEAD = crypto.SALT_LEN + crypto.NONCE_LEN + 16     # salt + nonce + tag GCM = 44 B


def rand_msg(n, seed=7):
    return np.random.default_rng(seed).bytes(n)


def load(name):
    return core.load_image_rgb(os.path.join(IMG_DIR, name + ".png"))


def write_sheet(wb, title, header, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="3465E0")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    for i, col in enumerate(ws.columns):
        w = (widths[i] if widths else None) or max(len(str(c.value or "")) for c in col) + 3
        ws.column_dimensions[col[0].column_letter].width = min(max(w, 10), 46)
    ws.freeze_panes = "A2"
    return ws


# ---------------------------------------------------------------------------
# m-bit LSB: trade-off kapasitas vs PSNR
# ---------------------------------------------------------------------------
def uji_mbit():
    print("[1/3] m-bit LSB ...")
    rows_fixed, rows_full = [], []
    for name in IMAGES:
        img = load(name)
        for m in range(1, 7):
            cap = core.capacity_bytes(img, m)
            for label, n, dest in (("tetap 5 KB", 5120, rows_fixed),
                                   ("penuh (kapasitas maks.)", cap - OVERHEAD, rows_full)):
                msg = rand_msg(n)
                st = core.embed(img, msg, PASSWORD, STEGO_KEY, m)
                ok = core.extract(st, PASSWORD, STEGO_KEY) == msg
                ps = metrics.psnr(img, st)
                dest.append([name + ".png", m, label, n, cap,
                             round(metrics.mse(img, st), 5), round(ps, 2),
                             round(metrics.changed_pixels_percent(img, st), 3),
                             "YA" if ps >= 30 else "TIDAK", "YA" if ok else "TIDAK"])
        print("   ", name)
    hdr = ["Citra", "m (bit/kanal)", "Skenario", "Pesan (byte)", "Kapasitas (byte)",
           "MSE", "PSNR (dB)", "Kanal Berubah (%)", "PSNR >= 30 dB", "Ekstraksi Utuh"]

    # grafik
    fig, ax = plt.subplots(1, 3, figsize=(15, 4.4))
    ms = np.arange(1, 7)
    theory = 10 * np.log10(255 ** 2 / ((4.0 ** ms - 1) / 6))
    for name in IMAGES:
        full = [r for r in rows_full if r[0] == name + ".png"]
        fixed = [r for r in rows_fixed if r[0] == name + ".png"]
        ax[0].plot(ms, [r[6] for r in full], "o-", label=name)
        ax[1].plot(ms, [r[6] for r in fixed], "o-", label=name)
    ax[0].plot(ms, theory, "k--", lw=2, label="teori (bit acak)")
    for a, t in ((ax[0], "Pesan memenuhi kapasitas maks."), (ax[1], "Pesan tetap 5 KB")):
        a.axhline(30, color="red", ls=":", label="ambang 30 dB")
        a.set_title(t); a.set_xlabel("m (bit per kanal)"); a.set_ylabel("PSNR (dB)")
        a.grid(alpha=.3)
    ax[0].legend(fontsize=7)
    cap512 = [core.capacity_bytes(load(IMAGES[0]), m) / 1024 for m in ms]
    ax[2].bar(ms, cap512, color="#3465e0")
    for x, v in zip(ms, cap512):
        ax[2].text(x, v, f"{v:.0f}", ha="center", va="bottom", fontsize=8)
    ax[2].set_title("Kapasitas citra 512x512 (KB)")
    ax[2].set_xlabel("m (bit per kanal)"); ax[2].set_ylabel("KB")
    fig.suptitle("Trade-off kapasitas vs kualitas pada m-bit LSB", fontweight="bold")
    fig.tight_layout()
    fig.savefig(os.path.join(OUT_DIR, "mbit_tradeoff.png"), dpi=130)
    plt.close(fig)
    return hdr, rows_fixed, rows_full


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    t0 = time.time()
    hdr, fixed, full = uji_mbit()
    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "mbit - pesan tetap 5KB", hdr, fixed)
    write_sheet(wb, "mbit - kapasitas penuh", hdr, full)
    wb.save(os.path.join(OUT_DIR, "hasil_mbit.xlsx"))
    print(f"Selesai dalam {time.time() - t0:.0f} detik -> {OUT_DIR}")


if __name__ == "__main__":
    main()
