import os
import numpy as np
from PIL import Image, ImageDraw
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

from stego import core, crypto, metrics

HERE = os.path.dirname(os.path.abspath(__file__))
IMG_DIR = os.path.join(HERE, "sample_images")
OUT_DIR = os.path.join(HERE, "hasil_uji")
os.makedirs(IMG_DIR, exist_ok=True)
os.makedirs(OUT_DIR, exist_ok=True)

PASSWORD = "sandi-rahasia-2026"
STEGO_KEY = "kunci-posisi-xyz"

MSG_SIZES = {"kecil (100 B)": 100, "sedang (1 KB)": 1024, "besar (5 KB)": 5120}

def buat_citra_uji():
    specs = []
    size = 512

    x = np.linspace(0, 255, size, dtype=np.uint8)
    grad = np.stack([np.tile(x, (size, 1)),
                     np.tile(x[::-1], (size, 1)),
                     np.tile(x.reshape(-1, 1), (1, size))], axis=-1).astype(np.uint8)
    specs.append(("01_gradien.png", grad))

    rng = np.random.default_rng(42)
    noise = rng.integers(0, 256, (size, size, 3), dtype=np.uint8)
    specs.append(("02_derau.png", noise))

    im = Image.new("RGB", (size, size), (30, 30, 60))
    d = ImageDraw.Draw(im)
    d.ellipse([60, 60, 300, 300], fill=(220, 80, 80))
    d.rectangle([250, 250, 460, 460], fill=(80, 200, 120))
    d.polygon([(400, 40), (480, 200), (320, 200)], fill=(240, 210, 60))
    specs.append(("03_geometri.png", np.array(im, dtype=np.uint8)))

    base = np.stack([
        np.tile(np.linspace(40, 210, size, dtype=np.float64), (size, 1)),
        np.tile(np.linspace(200, 30, size, dtype=np.float64).reshape(-1, 1), (1, size)),
        np.full((size, size), 120.0)], axis=-1)
    yy, xx = np.mgrid[0:size, 0:size]
    r = np.sqrt((xx - 256) ** 2 + (yy - 256) ** 2)
    base[..., 0] += 40 * np.cos(r / 30)
    base += rng.normal(0, 6, base.shape)
    foto = np.clip(base, 0, 255).astype(np.uint8)
    specs.append(("04_foto_sintetis.png", foto))

    board = np.zeros((size, size, 3), dtype=np.uint8)
    blk = 32
    for i in range(0, size, blk):
        for j in range(0, size, blk):
            if (i // blk + j // blk) % 2 == 0:
                board[i:i+blk, j:j+blk] = (200, 200, 200)
    imb = Image.fromarray(board)
    ImageDraw.Draw(imb).text((40, 240), "STEGO UJI", fill=(255, 0, 0))
    specs.append(("05_papan_teks.png", np.array(imb, dtype=np.uint8)))

    paths = []
    for name, arr in specs:
        p = os.path.join(IMG_DIR, name)
        if not os.path.exists(p):
            core.save_image(arr, p)
        paths.append(p)
    return paths


def uji_psnr_mse(paths):
    rows = []
    for p in paths:
        cover = core.load_image_rgb(p)
        cap = core.capacity_bytes(cover)
        for label, n in MSG_SIZES.items():
            pesan = os.urandom(n)
            stego = core.embed(cover, pesan, PASSWORD, STEGO_KEY)
            ps = metrics.psnr(cover, stego)
            ms = metrics.mse(cover, stego)
            chg = metrics.changed_pixels_percent(cover, stego)

            hasil = core.extract(stego, PASSWORD, STEGO_KEY)
            utuh = (hasil == pesan)
            rows.append({
                "citra": os.path.basename(p),
                "ukuran_pesan": label,
                "bytes": n,
                "kapasitas": cap,
                "MSE": round(ms, 5),
                "PSNR_dB": round(ps, 2),
                "kanal_berubah_%": round(chg, 4),
                "ekstraksi_utuh": "YA" if utuh else "TIDAK",
            })
    return rows


def uji_kerapuhan(paths):
    rows = []
    p = paths[0]
    cover = core.load_image_rgb(p)
    stego = core.embed(cover, b"pesan uji kerapuhan JPEG", PASSWORD, STEGO_KEY)
    png_path = os.path.join(OUT_DIR, "stego_uji_kerapuhan.png")
    core.save_image(stego, png_path)
    try:
        core.extract_from_file(png_path, PASSWORD, STEGO_KEY)
        png_ok = "BERHASIL"
    except Exception:
        png_ok = "GAGAL"
    jpg_path = os.path.join(OUT_DIR, "stego_uji_kerapuhan.jpg")
    Image.fromarray(stego).save(jpg_path, "JPEG", quality=90)
    try:
        core.extract_from_file(jpg_path, PASSWORD, STEGO_KEY)
        jpg_ok = "BERHASIL"
    except Exception:
        jpg_ok = "GAGAL (sesuai harapan: LSB rusak oleh kompresi JPEG)"
    rows.append({"format": "PNG (lossless)", "ekstraksi": png_ok})
    rows.append({"format": "JPEG q=90 (lossy)", "ekstraksi": jpg_ok})
    return rows


def buat_histogram(paths):
    for p in paths[:2]:  
        cover = core.load_image_rgb(p)
        stego = core.embed(cover, os.urandom(2048), PASSWORD, STEGO_KEY)
        hc = metrics.histogram(cover)
        hs = metrics.histogram(stego)
        fig, ax = plt.subplots(1, 2, figsize=(10, 3.2))
        ax[0].bar(range(256), hc, width=1, color="#3b6"); ax[0].set_title("Histogram Cover")
        ax[1].bar(range(256), hs, width=1, color="#36b"); ax[1].set_title("Histogram Stego")
        for a in ax:
            a.set_xlabel("Nilai byte (0-255)"); a.set_ylabel("Frekuensi")
        fig.suptitle(f"Perbandingan Histogram — {os.path.basename(p)}")
        fig.tight_layout()
        fig.savefig(os.path.join(OUT_DIR, f"histogram_{os.path.basename(p)}"), dpi=110)
        plt.close(fig)


def buat_lsb_plane(paths):
    for p in paths[:2]:
        cover = core.load_image_rgb(p)
        stego = core.embed(cover, os.urandom(3000), PASSWORD, STEGO_KEY)
        fig, ax = plt.subplots(1, 2, figsize=(9, 4.6))
        ax[0].imshow(metrics.lsb_plane(cover)); ax[0].set_title("Bidang LSB — Cover"); ax[0].axis("off")
        ax[1].imshow(metrics.lsb_plane(stego)); ax[1].set_title("Bidang LSB — Stego"); ax[1].axis("off")
        fig.suptitle(f"Steganalisis Visual (Enhanced LSB) — {os.path.basename(p)}")
        fig.tight_layout()
        fig.savefig(os.path.join(OUT_DIR, f"lsb_plane_{os.path.basename(p)}"), dpi=110)
        plt.close(fig)

def buat_berdampingan(paths):
    for p in paths[:2]:
        cover = core.load_image_rgb(p)
        stego = core.embed(cover, os.urandom(2048), PASSWORD, STEGO_KEY)
        ps = metrics.psnr(cover, stego)
        fig, ax = plt.subplots(1, 2, figsize=(9, 4.8))
        ax[0].imshow(cover); ax[0].set_title("Cover (asli)"); ax[0].axis("off")
        ax[1].imshow(stego); ax[1].set_title(f"Stego (PSNR {ps:.2f} dB)"); ax[1].axis("off")
        fig.suptitle(f"Cover vs Stego — {os.path.basename(p)}")
        fig.tight_layout()
        fig.savefig(os.path.join(OUT_DIR, f"berdampingan_{os.path.basename(p)}"), dpi=110)
        plt.close(fig)


def tulis_xlsx(rows_psnr, rows_rapuh):
    wb = Workbook()
    ws = wb.active
    ws.title = "PSNR-MSE"
    headers = ["Citra", "Ukuran Pesan", "Bytes", "Kapasitas (byte)",
               "MSE", "PSNR (dB)", "Kanal Berubah (%)", "Ekstraksi Utuh"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True); c.alignment = Alignment(horizontal="center")
    for r in rows_psnr:
        ws.append([r["citra"], r["ukuran_pesan"], r["bytes"], r["kapasitas"],
                   r["MSE"], r["PSNR_dB"], r["kanal_berubah_%"], r["ekstraksi_utuh"]])
    for col, w in zip("ABCDEFGH", [20, 16, 8, 15, 10, 10, 16, 14]):
        ws.column_dimensions[col].width = w

    ws2 = wb.create_sheet("Uji Kerapuhan JPEG")
    ws2.append(["Format", "Hasil Ekstraksi"])
    for c in ws2[1]:
        c.font = Font(bold=True)
    for r in rows_rapuh:
        ws2.append([r["format"], r["ekstraksi"]])
    ws2.column_dimensions["A"].width = 22
    ws2.column_dimensions["B"].width = 48

    out = os.path.join(OUT_DIR, "hasil_psnr_mse.xlsx")
    wb.save(out)
    return out


def main():
    print("1. Membuat citra uji...")
    paths = buat_citra_uji()
    print(f"   {len(paths)} citra siap di sample_images/")

    print("2. Uji PSNR & MSE (5 citra x 3 ukuran pesan)...")
    rows = uji_psnr_mse(paths)
    for r in rows:
        print(f"   {r['citra']:<22} {r['ukuran_pesan']:<14} "
              f"PSNR={r['PSNR_dB']:>7} dB  MSE={r['MSE']:>8}  utuh={r['ekstraksi_utuh']}")

    print("3. Uji kerapuhan (PNG vs JPEG)...")
    rapuh = uji_kerapuhan(paths)
    for r in rapuh:
        print(f"   {r['format']:<20} -> {r['ekstraksi']}")

    print("4. Membuat grafik histogram, bidang LSB, dan citra berdampingan...")
    buat_histogram(paths)
    buat_lsb_plane(paths)
    buat_berdampingan(paths)

    print("5. Menulis tabel hasil ke XLSX...")
    xlsx = tulis_xlsx(rows, rapuh)

    print("\nSelesai. Semua hasil ada di folder hasil_uji/")
    print("Tabel:", os.path.basename(xlsx))

if __name__ == "__main__":
    main()
